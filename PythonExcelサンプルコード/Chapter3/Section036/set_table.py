from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

wb = load_workbook('売上実績.xlsx')
ws = wb.active

table = Table(displayName='Table1', ref='B2:F12')
table_style = TableStyleInfo(name='TableStyleMedium9',
                             showRowStripes=True)

table.tableStyleInfo = table_style
ws.add_table(table)

wb.save('売上実績_変更後.xlsx')
