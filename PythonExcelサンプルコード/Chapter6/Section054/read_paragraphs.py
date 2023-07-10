from docx import Document
from openpyxl import Workbook

wb = Workbook()
ws = wb.active

document = Document('論文.docx')
paragraphs = document.paragraphs

for i, paragraph in enumerate(paragraphs):
    ws.cell(i + 1, 1).value = paragraph.text

wb.save('論文.xlsx')
