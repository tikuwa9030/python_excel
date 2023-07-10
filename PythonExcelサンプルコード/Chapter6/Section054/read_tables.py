from docx import Document
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.column_dimensions['A'].width = 40
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 15

document = Document('指摘一覧.docx')
tables = document.tables

for table in tables:
    for row in table.rows:
        row_list = []
        for cell in row.cells:
            text = cell.text
            row_list.append(text)
        ws.append(row_list)

wb.save('指摘一覧.xlsx')
