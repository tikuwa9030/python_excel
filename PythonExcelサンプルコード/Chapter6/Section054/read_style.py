from docx import Document
from openpyxl import Workbook

wb = Workbook()
ws = wb.active

document = Document('論文.docx')
paragraphs = document.paragraphs

heading1_style = 'Heading 1'

count = 0
for paragraph in paragraphs:
    if paragraph.style.name == heading1_style:
        count += 1
        ws.cell(count, 1).value = paragraph.text

wb.save('論文見出し.xlsx')
